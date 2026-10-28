""" Saxo importer

    SaxoTraderGO -> Account -> Historic reports -> Transactions -> Export -> Excel.
    Only the "Trades" sheet is used for trades:

        Account ID, Trade ID, Corporate Action Id, Event, Order ID, Adjusted Trade Date, Trade execution date,
        Trade Event Type, Open/Close, Traded Quantity, Price, Traded Value, Instrument, Instrument Symbol,
        Instrument ISIN, Instrument currency, Type, ...

    Instrument Symbol identifies a single contract (e.g. BTU/13X26C26:xcbf for options, BTU:xnys for stocks).
    Traded Value already includes the contract multiplier. Open/Close is used to verify that no history is missing.

    Stock splits on options are exported as a corporate action closing the old contract and opening the adjusted
    contract under the old symbol, while later trades use the new symbol (APH/16V26P140 -> APH/16V26P70 after 2:1).
    The split is not a disposal, so earlier trades are adjusted to the new contract and the corporate action rows
    are dropped.
"""

import datetime
import re
import sys

from importers import contracts

assetCategories = {
    "Stock": "STK",
    "Etf": "STK",
    "Etc": "STK",
    "Etn": "STK",
    "Bond": "STK",
    "StockOption": "OPT",
    "StockIndexOption": "OPT",
    "FuturesOption": "FOP",
    "ContractFutures": "FUT",
    "CfdOnStock": "CFD",
    "CfdOnEtf": "CFD",
    "CfdOnEtc": "CFD",
    "CfdOnEtn": "CFD",
    "CfdOnIndex": "CFD",
    "CfdOnFutures": "CFD",
    "CfdIndexOption": "CFD",
    "FxSpot": "CASH",
}


def readSheet(filename, sheetName):
    try:
        import openpyxl
    except ImportError:
        sys.exit("Reading Saxo reports requires openpyxl: pip install openpyxl")
    workbook = openpyxl.load_workbook(filename, read_only=True, data_only=True)
    if sheetName not in workbook.sheetnames:
        return None
    rows = workbook[sheetName].iter_rows(values_only=True)
    header = [str(h).replace("\xa0", " ").strip() for h in next(rows)]
    return [dict(zip(header, row)) for row in rows if any(v not in (None, "") for v in row)]


def isSaxoFile(filename):
    if not filename.lower().endswith(".xlsx"):
        return False
    try:
        trades = readSheet(filename, "Trades")
    except Exception:
        return False
    return trades is not None and (len(trades) == 0 or "Instrument Symbol" in trades[0])


def getContractKey(symbol, isin, assetCategory, multipliers):
    if assetCategory in ("OPT", "FOP"):
        parsed = contracts.parseSaxoSymbol(symbol, multipliers.get(symbol))
        if parsed is not None:
            return parsed["key"]
    if assetCategory == "STK":
        return "STK:" + (isin or symbol.split(":")[0])
    return "saxo:" + symbol


def parseTrade(record, multipliers):
    time = record["Trade execution date"]
    if not isinstance(time, datetime.datetime):
        time = datetime.datetime.fromisoformat(str(time))
    quantity = float(record["Traded Quantity"])
    price = abs(float(record["Traded Value"] or 0)) / abs(quantity) if quantity else 0
    assetCategory = assetCategories.get(record["Type"], record["Type"])

    trade = {
        "broker": "saxo",
        "accountId": record["Account ID"],
        "transactionID": "saxo:" + str(record["Trade ID"]),
        "brokerTransactionID": str(record["Trade ID"]),
        "ibOrderID": "saxo:" + str(record["Order ID"]),
        "contractKey": getContractKey(record["Instrument Symbol"], record.get("Instrument ISIN"), assetCategory, multipliers),
        "symbol": record["Instrument Symbol"],
        "description": record["Instrument"],
        "currency": record["Instrument currency"],
        "assetCategory": assetCategory,
        "tradePrice": price,
        "quantity": quantity,
        "tradeDate": time.strftime("%Y%m%d"),
        "tradeTime": time.strftime("%H%M%S"),
        "rawTime": str(time),
        "originalPrice": float(record.get("Price") or 0),
        "originalQuantity": quantity,
        "originalValue": float(record["Traded Value"] or 0),
    }
    if record["Instrument Symbol"] in multipliers:
        trade["multiplier"] = multipliers[record["Instrument Symbol"]]
    parsed = contracts.parseSaxoSymbol(record["Instrument Symbol"]) if assetCategory in ("OPT", "FOP") else None
    if parsed is not None:
        trade["expiryDate"] = parsed["expiryDate"]
    if record.get("Instrument ISIN"):
        trade["isin"] = record["Instrument ISIN"]
    if record.get("Open/Close") == "To Open":
        trade["openCloseHint"] = "O"
    elif record.get("Open/Close") == "To Close":
        trade["openCloseHint"] = "C"
    return trade


def splitSymbol(symbol, multiplier):
    """ Option strike is divided by the split multiplier: APH/16V26P140:xcbf with 2:1 -> APH/16V26P70:xcbf """
    match = re.match(r"^(.+/\d{1,2}[A-Z]\d{2}[CP])([0-9.]+)(:.*)$", symbol)
    if match is None:
        return symbol
    return match.group(1) + "%g" % (float(match.group(2)) / multiplier) + match.group(3)


def applyCorporateActions(corporateActions, trades, multipliers, audit):
    actions = {}
    for record in corporateActions:
        actions.setdefault(record["Corporate Action Id"], []).append(record)

    for actionId, records in actions.items():
        events = {r["Event"] for r in records}
        closes = [r for r in records if r.get("Open/Close") == "To Close"]
        opens = [r for r in records if r.get("Open/Close") == "To Open"]
        if events != {"Stock split"} or len(closes) != 1 or len(opens) != 1:
            """ Unknown corporate action: keep its rows as ordinary trades, so no position change is lost """
            for record in records:
                trade = parseTrade(record, multipliers)
                trades.append(trade)
            audit.add(
                "Saxo corporate action %s (%s) is not recognized, its %d rows were imported as ordinary trades "
                "(a close and reopen is then reported as a disposal)."
                % (actionId, ", ".join(sorted(str(e) for e in events)), len(records))
            )
            continue
        close = parseTrade(closes[0], multipliers)
        multiplier = abs(float(opens[0]["Traded Quantity"])) / abs(float(closes[0]["Traded Quantity"]))
        newSymbol = splitSymbol(close["symbol"], multiplier)
        """ A split changes the strike, not the contract multiplier """
        if close["symbol"] in multipliers:
            multipliers.setdefault(newSymbol, multipliers[close["symbol"]])
        splitTime = close["tradeDate"] + close["tradeTime"]
        for trade in trades:
            if trade["contractKey"] == close["contractKey"] and trade["tradeDate"] + trade["tradeTime"] < splitTime:
                trade["quantity"] *= multiplier
                trade["tradePrice"] /= multiplier
                if newSymbol != close["symbol"]:
                    trade["symbol"] = newSymbol
                    trade["contractKey"] = getContractKey(newSymbol, None, trade["assetCategory"], multipliers)
                    trade["description"] = trade["description"] + " (adjusted for split)"
        audit.add("Saxo %s: stock split with ratio %g on %s, earlier trades adjusted to %s (not a disposal)."
                  % (close["symbol"], multiplier, close["tradeDate"], newSymbol))


def importTrades(filenames, audit):
    records = {}
    for filename in filenames:
        for record in readSheet(filename, "Trades") or []:
            """ Overlapping reports, keep each execution once """
            records[(record.get("Corporate Action Id"), str(record["Trade ID"]))] = record
    records = [records[key] for key in sorted(records, key=lambda k: (str(k[0]), k[1]))]

    multipliers, conflicts = contracts.inferMultipliers(
        (r["Instrument Symbol"], float(r["Traded Value"] or 0), float(r.get("Price") or 0), float(r["Traded Quantity"]))
        for r in records)
    for symbol, ratios in sorted(conflicts.items()):
        audit.add("Saxo %s: executions imply different multipliers %s, the multiplier is left unknown in the "
                  "contract key." % (symbol, ratios))

    trades = [parseTrade(r, multipliers) for r in records if not r.get("Corporate Action Id")]
    applyCorporateActions([r for r in records if r.get("Corporate Action Id")], trades, multipliers, audit)
    return trades


def importCashTransactions(filenames):
    """ Non-trade rows of the Transactions sheet (financing and carrying costs, securities lending fees, transfers,
        dividends). Amount is in the account currency (Booked Amount and _Currency). """
    rows = {}
    for filename in filenames:
        for record in readSheet(filename, "Transactions") or []:
            if record.get("Transaction Type") in ("Trade", None):
                continue
            key = record.get("Bk Record Id") or record.get("Booking Id") or (record.get("Corporate Action Id"), record.get("Position ID"))
            rows[key] = record

    cash = []
    for key, record in rows.items():
        date = record["Trade Date"]
        if not isinstance(date, datetime.datetime):
            date = datetime.datetime.fromisoformat(str(date))
        event = str(record.get("Event") or "")
        transactionType = str(record.get("Transaction Type") or "")
        if transactionType == "Cash Transfer":
            category = "funding"
        elif event in ("Financing Cost", "Carrying Cost"):
            category = "financing"
        elif event == "Securities Lending Client Fee":
            category = "securities lending income"
        elif "dividend" in event.lower():
            category = "dividend"
        elif "tax" in event.lower():
            category = "withholding tax"
        elif "interest" in event.lower():
            category = "interest"
        elif transactionType == "Corporate action":
            category = "corporate action"
        else:
            category = "unknown"
        cash.append({
            "broker": "saxo",
            "accountId": record["Account ID"],
            "transactionID": "saxo:%s" % (key,),
            "operation": "%s / %s" % (transactionType, event),
            "category": category,
            "symbol": record.get("Instrument Symbol") or None,
            "isin": record.get("Instrument ISIN") or None,
            "date": date.strftime("%Y%m%d"),
            "time": "000000",
            "amount": float(record.get("Booked Amount") or 0),
            "currency": record.get("_Currency") or record.get("Currency"),
            "brokerEUR": float(record.get("Amount") or 0) if record.get("Currency") == "EUR" else None,
            "comment": str(record.get("Comment") or ""),
        })
    cash.sort(key=lambda c: (c["date"], c["transactionID"]))
    return cash
